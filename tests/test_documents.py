import io
from dataclasses import replace
import subprocess
import zipfile
import pytest
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pptx import Presentation
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from deepdive.api import create_app
from deepdive.config import Settings
from deepdive.document_parser import extract
from deepdive.documents import Documents, chunks
from deepdive.guardrails import scope_for
from deepdive.models import RunRequest, ReviewRequest
from deepdive.ops import backup, restore
from deepdive.service import ResearchService, Conflict
from deepdive.storage import Store
from .fakes import FixtureGateway, FixtureHTTP


class DocumentGateway(FixtureGateway):
    """Synthetic vectors/ranks for contract tests; not a semantic-quality benchmark."""
    def __init__(self):
        super().__init__()
        self.embedded = []

    def embed(self, context, texts):
        context.reserve()
        self.embedded.extend(texts)
        return [[1., float(any(w in text.lower() for w in ["automobile", "vehicle", "car"])),
                 float("follow-up" in text.lower())] for text in texts]

    def ask(self, context, role, instruction, payload, schema):
        if role == "document reranker":
            context.reserve()
            self.roles.append(role)
            return schema.model_validate({"passage_ids": [p["id"] for p in payload["passages"]]})
        return super().ask(context, role, instruction, payload, schema)


def pdf_bytes(text="Study enrolled 20 adults in 2025.", encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 20 700 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = content
    if encrypted:
        writer.encrypt("private")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def test_parser_formats_and_precise_chunk_locations():
    parsed = extract(pdf_bytes(), ".pdf")
    assert parsed["blocks"][0]["page"] == 1
    assert "20 adults" in parsed["blocks"][0]["text"]
    assert list(chunks(parsed["blocks"]))[0]["location"].startswith("page 1")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:body><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Methods</w:t></w:r></w:p>'
            '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Adults</w:t></w:r></w:p></w:tc>'
            '<w:tc><w:p><w:r><w:t>20</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>')
    docx = extract(out.getvalue(), ".docx")
    assert docx["blocks"][0]["section"] == "Methods"
    assert "Adults | 20" in docx["blocks"][0]["text"]
    parsed = extract(("# One\n\n" + "A sentence about naïve café research. " * 200 + "\n# Two\n\nFinal section.").encode(), ".md")
    windows = list(chunks(parsed["blocks"]))
    assert len(windows) > 2
    assert all(0 < w["tokens"] <= 384 for w in windows)
    for w in windows:
        assert w["text"] == parsed["blocks"][w["block"]]["text"][w["start"]:w["end"]]
        assert "�" not in w["text"]
    assert windows[-1]["section"] == "Two"
    assert windows[1]["start"] < windows[0]["end"]


@pytest.mark.parametrize("data,suffix", [(b"fake", ".pdf"), (b"\x00binary", ".txt"), (b"\xff", ".txt"),
                                        (b"   ", ".txt"), (b"content", ".exe")])
def test_invalid_documents_fail(data, suffix):
    with pytest.raises((ValueError, UnicodeError)):
        extract(data, suffix)


def test_encrypted_and_scanned_pdf_fail():
    with pytest.raises(ValueError, match="Encrypted"):
        extract(pdf_bytes(encrypted=True), ".pdf")
    with pytest.raises(ValueError, match="OCR"):
        extract(pdf_bytes(""), ".pdf")


def test_docx_xml_entity_and_zip_limits():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("word/document.xml", '<!DOCTYPE x [<!ENTITY y SYSTEM "file:///etc/passwd">]><x>&y;</x>')
    with pytest.raises(Exception, match="EntitiesForbidden"):
        extract(out.getvalue(), ".docx")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 4_000_001)
    with pytest.raises(ValueError, match="4 MB"):
        extract(out.getvalue(), ".docx")


def test_atomic_dedup_scope_removal_and_parser_timeout(tmp_path, monkeypatch):
    store = Store(tmp_path)
    docs = Documents(store)
    doc = docs.ingest("session:a", "study.txt", b"Study enrolled 20 adults.")
    assert docs.ingest("session:a", "renamed.txt", b"Study enrolled 20 adults.")["id"] == doc["id"]
    other = docs.ingest("session:b", "study.txt", b"Study enrolled 20 adults.")
    assert other["id"] != doc["id"]
    with pytest.raises(ValueError, match="outside"):
        docs.selected("session:b", [doc["id"]])
    with pytest.raises(ValueError, match="filename"):
        docs.ingest("session:a", "../study.txt", b"text")
    with pytest.raises(ValueError, match="5 MB"):
        docs.ingest("session:a", "large.txt", b"x" * 5_000_001)
    docs.remove("session:a", doc["id"])
    assert docs.list("session:a") == []
    assert store.get("source", doc["source_id"], "session:a")
    assert docs.ingest("session:a", "study.txt", b"Study enrolled 20 adults.")["id"] == doc["id"]
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("parser", 30)
    monkeypatch.setattr("deepdive.documents.subprocess.run", timeout)
    with pytest.raises(ValueError, match="30 seconds"):
        docs.ingest("session:a", "timeout.txt", b"different document")
    assert len(docs.list("session:a")) == 1


def test_upload_api_boundaries_and_document_workflow_restart(tmp_path):
    settings = Settings(tmp_path / "original", api_token="test-token")
    gateway, http = DocumentGateway(), FixtureHTTP()
    s = ResearchService(settings, gateway=gateway, http=http)
    try:
        session = s.create_session("Documents")
        scope = scope_for(session)
        with TestClient(create_app(s, settings)) as api:
            assert api.post("/documents", content=b"text").status_code == 401
            api.headers["Authorization"] = "Bearer test-token"
            response = api.post("/documents", params={"session_id": session["id"], "filename": "study.txt"},
                               content=b"Study enrolled 20 adults in 2025.\nFollow-up lasted 12 months.",
                               headers={"Content-Type": "application/octet-stream"})
            assert response.status_code == 201, response.text
            doc = response.json()
            assert api.get("/documents", params={"session_id": session["id"]}).json()["documents"][0]["id"] == doc["id"]
            assert api.post("/documents", params={"session_id": session["id"], "filename": "big.txt"},
                            content=b"x" * 5_000_001).status_code == 413
            assert api.post("/runs", content=b"x" * 200001).status_code == 413
        run = s.create_run(RunRequest(question="Study participants and follow-up", session_id=session["id"],
                          document_ids=[doc["id"]], research_mode="documents"), enqueue=False)
        s.execute(run["id"])
        assert s.get_run(run["id"])["status"] == "waiting_review", s.get_run(run["id"])
        assert http.calls == []
        assert "document reranker" in gateway.roles
        with pytest.raises(Conflict):
            s.remove_document(scope, doc["id"])
        s.close()
        archive = tmp_path / "research.zip"
        backup(settings.data, archive)
        restored = tmp_path / "restored"
        restore(archive, restored)
        settings = replace(settings, data=restored)
        s = ResearchService(settings, gateway=DocumentGateway(), http=http)
        s.review(run["id"], ReviewRequest(decision="revise", feedback="Emphasize follow-up duration"), enqueue=False)
        s.execute(run["id"])
        assert s.get_run(run["id"])["status"] == "waiting_review"
        assert s.get_run(run["id"])["version"] == 2
        s.review(run["id"], ReviewRequest(decision="approve"), enqueue=False)
        s.execute(run["id"])
        assert s.get_run(run["id"])["status"] == "completed"
        report = s.artifacts.render(run["id"])
        assert "study.txt" in report[1] and "Location: section 1" in report[1]
        artifact = next(a for a in s.export_run(run["id"])["artifacts"] if a["name"] == "slides.pptx")
        slides = Presentation(s.artifacts.path(run["id"], artifact["id"]))
        assert any("section 1" in shape.text for slide in slides.slides for shape in slide.shapes if shape.has_text_frame)
        assert http.calls == []
        s.remove_document(scope, doc["id"])
        assert s.documents.list(scope) == []
        assert "Study enrolled" in s.artifacts.render(run["id"])[1]
        s.delete_session(session["id"])
        with s.store.connection() as db:
            assert db.execute("SELECT count(*) FROM document_vectors").fetchone()[0] == 0
            assert db.execute("SELECT count(*) FROM document_search").fetchone()[0] == 0
        assert s.store.list("document", scope) == []
    finally:
        s.close()


def test_document_only_blocks_tools_and_cross_scope_selection(tmp_path):
    s = ResearchService(Settings(tmp_path), gateway=DocumentGateway(), http=FixtureHTTP())
    try:
        session = s.create_session("Private")
        doc = s.documents.ingest(scope_for(session), "private.txt", b"A private research document.")
        other = s.create_session("Other")
        with pytest.raises(ValueError, match="outside"):
            s.create_run(RunRequest(question="Private", session_id=other["id"], document_ids=[doc["id"]], research_mode="documents"))
        run = s.create_run(RunRequest(question="Private", session_id=session["id"], document_ids=[doc["id"]], research_mode="documents"), enqueue=False)
        _, _, tools = s.runtime(run["id"])
        for name, args in [("search_web", {"query": "private"}), ("fetch_source", {"url": "https://example.org"}),
                           ("fetch_market_series", {"symbol": "MSFT"})]:
            assert name not in tools.schemas()
            assert "disabled" in tools.call("researcher", name, args)["error"]
        assert tools.http.calls == []
    finally:
        s.close()
