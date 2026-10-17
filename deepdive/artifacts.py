"""Provenance-preserving exports adapted from the original chart/deck features."""
import base64
import html
import json
import math
import re
import shutil
import textwrap
from datetime import datetime, timezone
from urllib.parse import urlsplit
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image
from .storage import digest

STYLE = """
body{margin:0;background:#141a22;color:#f5f7fb;font:17px/1.65 system-ui,sans-serif}
main{max-width:1050px;margin:3rem auto;padding:0 2rem}
h1,h2{font-family:Georgia,serif;font-weight:400}h1{font-size:2.8rem;line-height:1.15}
h2{color:#4ec9b0}a{color:#8de4cd}section{padding:1.5rem 0;border-bottom:1px solid #384452}
small,.muted{color:#bdc9d7}table{border-collapse:collapse;width:100%}
td,th{border:1px solid #384452;padding:.6rem;text-align:left}
img{max-width:100%;border-radius:12px}.tag{color:#f0a46b}
details{padding:.8rem;border:1px solid #384452;border-radius:8px;margin:.5rem 0}
.slide{min-height:75vh;padding:3rem 0;page-break-after:always}
@media print{body{background:white;color:#111}a,h2{color:#145b50}.slide{min-height:auto}}
"""


def escape(value):
    return html.escape(str(value), quote=True)


def safe_link(url):
    return url if urlsplit(url).scheme in {"http", "https"} else ""


def md_text(value):
    text = html.escape(str(value))
    for char in ["\\", "[", "]", "(", ")", "*", "_", "#", chr(96)]:
        text = text.replace(char, "\\" + char)
    return text


def utc_date(value):
    return datetime.fromtimestamp(value, timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


class Artifacts:
    def __init__(self, service):
        self.s = service
        self.root = service.settings.data / "artifacts"

    def directory(self, run_id, version=None):
        if not re.fullmatch(r"run_[a-f0-9]{20}", run_id):
            raise ValueError("Invalid artifact run ID")
        run = self.s.get_run(run_id)
        directory = self.root / run_id / f"v{version or run['version']}"
        directory.mkdir(parents=True, exist_ok=True)
        return directory

    def save(self, run, name, content, kind, claim_ids=None, source_ids=None):
        path = self.directory(run["id"]) / name
        temporary = path.with_suffix(path.suffix + ".tmp")
        if isinstance(content, str):
            temporary.write_text(content, encoding="utf-8")
        else:
            temporary.write_bytes(content)
        temporary.replace(path)
        record = {"id": "artifact_" + digest([run["id"], run["version"], name])[:24],
                  "run_id": run["id"], "version": run["version"], "name": name,
                  "path": str(path.relative_to(self.s.settings.data)), "kind": kind,
                  "claim_ids": claim_ids or [], "source_ids": source_ids or []}
        return self.s.store.put("artifact", record, run["scope"])

    def path(self, run_id, artifact_id):
        run = self.s.get_run(run_id)
        artifact = self.s.store.get("artifact", artifact_id, run["scope"])
        if not artifact or artifact["run_id"] != run_id:
            raise KeyError(artifact_id)
        root = (self.root / run_id).resolve()
        path = (self.s.settings.data / artifact["path"]).resolve()
        if not path.is_relative_to(root) or not path.is_file():
            raise KeyError(artifact_id)
        return path

    def charts(self, run_id):
        run = self.s.get_run(run_id)
        evidence = self.s.evidence(run_id)
        claims = [self.s.store.get("claim", c, run["scope"]) for c in run.get("published_claim_ids", [])]
        records = []
        rows = [[s["title"][:65], s["quality"]["label"], utc_date(s["retrieved_at"])] for s in evidence["sources"]]
        if rows:
            records.append(self._table_image(run, "evidence.png", ["Source", "Assessment", "Retrieved"], rows[:30],
                                             source_ids=[s["id"] for s in evidence["sources"]]))
        if claims:
            rows = [[c["id"][-8:], c["text"]] for c in claims]
            for start in range(0, len(rows), 4):
                name = "insights.png" if start == 0 else f"insights-{start // 4 + 1}.png"
                records.append(self._table_image(run, name, ["Claim", "Model-verified finding"], rows[start:start + 4],
                                                 claim_ids=[c["id"] for c in claims[start:start + 4]]))
        for source in evidence["sources"]:
            for table in source.get("tables", [])[:2]:
                rows = table["rows"]
                if len(rows) < 3 or len(rows[0]) < 2:
                    continue
                headers = [str(x) for x in rows[0]]
                points = []
                for row in rows[1:31]:
                    if len(row) != len(headers):
                        continue
                    raw = str(row[1]).strip()
                    if not re.fullmatch(r"-?\d+(?:,\d{3})*(?:\.\d+)?", raw):
                        continue
                    y = float(raw.replace(",", ""))
                    if math.isfinite(y):
                        points.append((str(row[0]), y))
                if len(points) < 2:
                    continue
                fig, ax = plt.subplots(figsize=(11, 5))
                labels, values = zip(*points)
                ax.plot(range(len(values)), values, color="#168b78", marker="o", markersize=3)
                step = max(1, len(labels) // 6)
                ax.set_xticks(list(range(0, len(labels), step)), labels[::step], rotation=25, ha="right")
                ax.set_xlabel(headers[0]); ax.set_ylabel(headers[1])
                ax.set_title(textwrap.fill(table.get("caption") or source["title"], 90), fontsize=11)
                fig.text(.02, .01, f"Source: {source['id']} | column 1 as captured; retrieved {source['retrieved_at']}", fontsize=7)
                fig.tight_layout(rect=(0, .05, 1, 1))
                name = f"data-{source['id'][-12:]}-{table['index']}.png"
                path = self.directory(run_id) / name
                temp = path.with_suffix(".tmp")
                fig.savefig(temp, format="png", dpi=140); plt.close(fig); temp.replace(path)
                records.append(self.save(run, name, path.read_bytes(), "chart", source_ids=[source["id"]]))
        return records

    def _table_image(self, run, name, headers, rows, claim_ids=None, source_ids=None):
        wrapped = [[textwrap.fill(str(v), 72 if col == len(headers) - 1 else 45) for col, v in enumerate(row)] for row in rows]
        height = min(40, max(3, 1 + sum(max(str(v).count("\n") + 1 for v in row) for row in wrapped) * .24))
        fig, ax = plt.subplots(figsize=(13, height))
        ax.axis("off")
        table = ax.table(cellText=wrapped, colLabels=headers, loc="center", cellLoc="left", bbox=[0, 0, 1, 1])
        table.auto_set_font_size(False); table.set_fontsize(8)
        for (row, col), cell in table.get_celld().items():
            cell.set_edgecolor("#cbd5dd")
            if row == 0:
                cell.set_facecolor("#143e39"); cell.set_text_props(color="white", weight="bold")
        path = self.directory(run["id"]) / name
        temp = path.with_suffix(".tmp")
        fig.savefig(temp, format="png", dpi=130, bbox_inches="tight"); plt.close(fig); temp.replace(path)
        return self.save(run, name, path.read_bytes(), "table", claim_ids, source_ids)

    def render(self, run_id):
        run = self.s.get_run(run_id)
        claims = {c: self.s.store.get("claim", c, run["scope"]) for c in run.get("published_claim_ids", [])}
        evidence = self.s.evidence(run_id)
        passages = {p["id"]: p for p in evidence["passages"]}
        sources = {s["id"]: s for s in evidence["sources"]}
        def bullet(cid):
            claim = claims[cid]
            citations = " ".join(f'<a href="#{escape(pid)}">[{escape(pid[-8:])}]</a>' for pid in claim["passage_ids"])
            return f'<li>{escape(claim["text"])} {citations}</li>'
        sections = run.get("outline", {}).get("sections", [])
        title = run.get("outline", {}).get("title", run["request"]["question"])
        body = f'<h1>{escape(title)}</h1><p class="tag">DeepDive · evidence-backed research</p>'
        body += '<p>Claims passed automated source-support checks. This is not a guarantee of truth. Coverage gaps are shown below.</p>'
        summary = run.get("summary_ids") or run.get("outline", {}).get("summary_claim_ids", [])
        if summary:
            body += '<section><h2>Executive findings</h2><ul>' + "".join(bullet(c) for c in summary if c in claims) + '</ul></section>'
        if not claims:
            body += '<p><strong>Insufficient evidence. No substantive claim passed the checks.</strong></p>'
        for section in sections:
            body += f'<section><h2>{escape(section["title"])}</h2><ul>' + "".join(bullet(c) for c in section["claim_ids"] if c in claims) + '</ul></section>'
        body += '<section><h2>Coverage and limitations</h2><ul>'
        for item in run.get("coverage", []):
            for gap in item["missing_facets"]:
                body += f'<li>Unanswered facet: {escape(gap)}</li>'
            for conflict in item["contradictions"]:
                body += f'<li>Unresolved disagreement (assessment, not a verified finding): {escape(conflict)}</li>'
        body += '</ul><p>Public discovery can be blocked. Source dates may be unavailable. Historical snapshots do not imply current accuracy.</p></section>'
        body += '<section><h2>Evidence passages and references</h2>'
        used = sorted({p for c in claims.values() for p in c["passage_ids"]})
        for pid in used:
            p = passages[pid]; source = sources[p["source_id"]]
            link = safe_link(source.get("final_url", ""))
            body += f'<details id="{escape(pid)}"><summary>{escape(pid)} — {escape(source["title"])}</summary>'
            body += f'<blockquote>{escape(p["text"])}</blockquote><p><a href="{escape(link)}" rel="noreferrer">Original source</a> · retrieved {utc_date(source["retrieved_at"])}</p></details>'
        body += '</section>'
        markdown = [f"# {md_text(title)}", "", "Automated source-support checks are fallible; see coverage and evidence.", ""]
        for section in sections:
            markdown += [f"## {md_text(section['title'])}", ""]
            for cid in section["claim_ids"]:
                if cid in claims:
                    c = claims[cid]
                    markdown.append("- " + md_text(c["text"]) + " " + " ".join(f"[{p}](#ref-{p})" for p in c["passage_ids"]))
            markdown.append("")
        markdown += ["## Coverage gaps", ""]
        for item in run.get("coverage", []):
            markdown.extend("- " + md_text(g) for g in item["missing_facets"] + item["contradictions"])
        markdown += ["", "## References", ""]
        for pid in used:
            source = sources[passages[pid]["source_id"]]
            markdown += [f"<a id=\"ref-{pid}\"></a>", f"- [{pid}] {md_text(source['title'])}: {safe_link(source.get('final_url', ''))}",
                         "  Retrieved: " + utc_date(source["retrieved_at"]), "  Passage: " + md_text(passages[pid]["text"])]
        return body, "\n".join(markdown), claims, evidence

    def publish(self, run_id, charts):
        run = self.s.get_run(run_id)
        body, markdown, claims, evidence = self.render(run_id)
        title = escape(run["request"]["question"])
        head = f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title><style>{STYLE}</style></head><body><main>'
        self.save(run, "report.html", head + body + "</main></body></html>", "report", list(claims))
        self.save(run, "report.md", markdown, "report", list(claims))
        slides = [f'<section class="slide"><h1>{title}</h1><p>DeepDive · research with a visible evidence trail</p></section>']
        presentation = Presentation()
        presentation.slide_width = Inches(13.33); presentation.slide_height = Inches(7.5)
        def add_slide(heading, texts):
            lines = []
            for text in texts:
                lines.extend(textwrap.wrap(text, width=95) + [""])
            for start in range(0, max(1, len(lines)), 16):
                slide = presentation.slides.add_slide(presentation.slide_layouts[1])
                slide.shapes.title.text = heading[:140] + (" (continued)" if start else "")
                slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
                frame = slide.placeholders[1].text_frame
                frame.clear()
                frame.word_wrap = True
                paragraph = frame.paragraphs[0]
                paragraph.text = "\n".join(lines[start:start + 16])
                paragraph.font.size = Pt(16)
        add_slide("DeepDive research", [run["request"]["question"], "Automated verification is fallible. Sources and coverage gaps accompany this report."])
        for section in run.get("outline", {}).get("sections", []):
            ids = [c for c in section["claim_ids"] if c in claims]
            for start in range(0, len(ids), 3):
                chunk = ids[start:start + 3]
                texts = ["[" + c + "] " + claims[c]["text"] + " [" + ", ".join(claims[c]["passage_ids"]) + "]" for c in chunk]
                slides.append('<section class="slide"><h2>' + escape(section["title"]) + '</h2><ul>' + "".join("<li>" + escape(t) + "</li>" for t in texts) + "</ul></section>")
                add_slide(section["title"], texts)
        gaps = [g for item in run.get("coverage", []) for g in item["missing_facets"]]
        add_slide("Coverage and limitations", gaps or ["No missing facets reported by the coverage assessor; independent review is still necessary."])
        for chart in charts:
            path = self.path(run_id, chart["id"])
            encoded = base64.b64encode(path.read_bytes()).decode()
            slides.append(f'<section class="slide"><h2>{escape(chart["name"])}</h2><img alt="{escape(chart["name"])}" src="data:image/png;base64,{encoded}"></section>')
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = chart["name"]
            with Image.open(path) as captured:
                width, height = captured.size
            scale = min(12 / width, 5.8 / height)
            slide.shapes.add_picture(str(path), Inches(.5), Inches(1.2), width=Inches(width * scale), height=Inches(height * scale))
        refs = [s["title"] + " — " + s.get("final_url", "") for s in evidence["sources"]]
        for start in range(0, len(refs), 5):
            add_slide("Sources", refs[start:start + 5])
        slides.append('<section class="slide"><h2>Evidence and limitations</h2>' + body[body.find('<section><h2>Coverage'): ] + '</section>')
        self.save(run, "slides.html", head + "".join(slides) + "</main></body></html>", "presentation", list(claims))
        path = self.directory(run_id) / "slides.pptx"
        temp = path.with_suffix(".tmp")
        presentation.save(temp); temp.replace(path)
        self.save(run, "slides.pptx", path.read_bytes(), "presentation", list(claims))
        self.save(run, "audit.json", json.dumps(self.s.export_run(run_id), indent=2, ensure_ascii=False), "audit", list(claims))

    def delete_run(self, run_id):
        root = self.root.resolve()
        path = (root / run_id).resolve()
        if not path.is_relative_to(root) or path == root:
            raise ValueError("Artifact deletion escaped its root")
        if path.exists():
            shutil.rmtree(path)
